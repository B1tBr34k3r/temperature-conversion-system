import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_report():
    doc = docx.Document()

    # Set page margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styling helper
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # XML Helper functions for table styling
    def set_cell_background(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep navy blue
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Primary blue
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        return p

    def add_para(text, bold_prefix=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.bold = True
            r_pre.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.bold = True
            r_pre.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        return p

    def add_callout(title, text, color_hex="EFF6FF", border_color="2563EB"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, color_hex)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        # Add left border
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
            f'<w:top w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(3)
        r1 = cp.add_run(f"📌 {title}: ")
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        r2 = cp.add_run(text)
        r2.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_image_figure(img_path, caption, width=Inches(6.2)):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            doc.add_picture(img_path, width=width)
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(caption)
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # =========================================================
    # DOCUMENT COVER / HEADER
    # =========================================================
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_title.add_run("TEMPERATURE CONVERSION SYSTEM")
    r.font.name = 'Calibri'
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Milestone 2: Core Module Implementation & Prototyping")
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    # Metadata Box Table
    meta_table = doc.add_table(rows=9, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Title:", "Temperature Conversion System (Web-Based)"),
        ("Milestone Phase:", "Milestone 2 (Core Module Implementation & Prototyping)"),
        ("Authors:", "Prakhar Harne and Naman Shrivastav"),
        ("GitHub Repository:", "https://github.com/B1tBr34k3r/temperature-conversion-system"),
        ("Live Web Prototype:", "https://b1tbr34k3r.github.io/temperature-conversion-system/"),
        ("Platform & Stack:", "HTML5, Modern CSS3, JavaScript (ES6+), Web APIs, LocalStorage"),
        ("Architecture Model:", "Client-Side MVC Architecture with High-Precision Scientific Engine"),
        ("Current Version:", "Version 2.0 (Milestone 2 Deliverable)"),
        ("Date of Submission:", "October 2026")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, 60, 60, 100, 100)
        set_cell_margins(cell_val, 60, 60, 100, 100)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Callout regarding Milestone 1 Evaluation Feedback Incorporation
    add_callout(
        "Milestone 1 Feedback Resolution",
        "This Milestone 2 report fully addresses all review feedback from Milestone 1: (1) UML diagrams (Use Case, Class, Sequence) are now visually rendered and embedded directly into the report; (2) A rigorous Risk Assessment analyzing browser compatibility, localStorage quotas, and IEEE 754 precision has been integrated; and (3) The working prototype web application has been implemented and validated."
    )

    # =========================================================
    # 1. EXECUTIVE SUMMARY & MILESTONE 2 SCOPE
    # =========================================================
    add_heading_1("1. Executive Summary & Milestone Objectives")
    add_para(
        "Building upon the foundational requirement analysis and system specifications established in Milestone 1, Milestone 2 delivers the core computational engine, robust validation mechanisms, and an interactive prototype of the web-based Temperature Conversion System. The project adheres strictly to standard web technologies (HTML5, CSS3, ES6+ JavaScript) executed client-side within modern browsers, avoiding external runtime dependencies such as desktop C++ or Python frameworks."
    )
    add_para(
        "The overarching objectives realized in Milestone 2 include:"
    )
    add_bullet(" Implementation of the bidirectional conversion algorithms across Celsius (°C), Fahrenheit (°F), and Kelvin (K) with high-precision IEEE 754 floating-point arithmetic handling.", "1. Core Conversion Engine:")
    add_bullet(" Implementation of multi-stage input validation covering data type parsing, NaN/Infinity checks, and strict enforcement of thermodynamic lower limits (Absolute Zero: -273.15°C / -459.67°F / 0 K).", "2. Thermodynamic Boundary Validation:")
    add_bullet(" Construction of an accessible, responsive, and intuitive graphical user interface supporting dynamic scale selection, instantaneous conversion, precision configuration, unit swapping, and formula explanations.", "3. Interactive Web Prototype:")
    add_bullet(" Inclusion of session-persistent conversion history utilizing the browser's localStorage API, coupled with an automated in-memory fallback to handle private browsing and quota restrictions.", "4. History Persistence & Resiliency:")
    add_bullet(" Execution of a comprehensive automated unit test suite verifying accuracy, symmetry, edge cases, and numerical boundary conditions.", "5. Quality Assurance & Verification:")

    # =========================================================
    # 2. CORE MATHEMATICAL MODEL & FLOATING-POINT PRECISION
    # =========================================================
    add_heading_1("2. Mathematical Formulation & High-Precision Floating-Point Handling")
    add_para(
        "Temperature conversion requires rigorous scientific transformation formulas. In a digital computing environment governed by the IEEE 754 standard for double-precision binary floating-point numbers (64-bit), fractions such as 9/5 (1.8) and 5/9 (0.555555...) cannot be stored with infinite decimal precision. Therefore, the implementation incorporates dedicated stabilization strategies."
    )
    
    add_heading_2("2.1 Scientific Conversion Formulas")
    add_para("The core computational module implements the following direct transformations:")
    
    # Formula Table
    formula_tbl = doc.add_table(rows=7, cols=3)
    formula_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_headers = ["Transformation", "Standard Mathematical Formula", "Computational Implementation"]
    f_rows = [
        ("Celsius to Fahrenheit", "F = (C × 9/5) + 32", "(c * 9.0 / 5.0) + 32.0"),
        ("Fahrenheit to Celsius", "C = (F - 32) × 5/9", "(f - 32.0) * (5.0 / 9.0)"),
        ("Celsius to Kelvin", "K = C + 273.15", "c + 273.15"),
        ("Kelvin to Celsius", "C = K - 273.15", "k - 273.15"),
        ("Fahrenheit to Kelvin", "K = ((F - 32) × 5/9) + 273.15", "((f - 32.0) * (5.0 / 9.0)) + 273.15"),
        ("Kelvin to Fahrenheit", "F = ((K - 273.15) × 9/5) + 32", "((k - 273.15) * 9.0 / 5.0) + 32.0")
    ]
    for c_idx, h_text in enumerate(f_headers):
        cell = formula_tbl.rows[0].cells[c_idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    for r_idx, (col0, col1, col2) in enumerate(f_rows):
        row = formula_tbl.rows[r_idx + 1]
        bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate([col0, col1, col2]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            if c_idx == 2:
                r.font.name = 'Consolas'
                r.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_heading_2("2.2 IEEE 754 Floating-Point Precision & Rounding Strategy")
    add_para(
        "Standard JavaScript numbers represent double-precision binary floating-point values according to IEEE 754. Because numbers like 0.1, 0.2, or repeating decimals such as 5/9 cannot be represented exactly in binary fractions, naïve arithmetic can yield precision drift (e.g., 0.1 + 0.2 producing 0.30000000000000004). To eliminate arithmetic drift and presentation anomalies, the system introduces two defensive controls:"
    )
    add_bullet(" In unit testing and mathematical symmetry verification, differences are validated within a tight numerical tolerance (epsilon = 1e-5 to 1e-9) rather than requiring absolute binary identity.", "Epsilon-Based Boundary Assertion:")
    add_bullet(" When formatting for display or history logging, the formatPrecision(val, p) method adds Number.EPSILON prior to rounding factor multiplication, preventing truncation errors at boundary digits.", "Exponent-Scaled Rounding:")

    # =========================================================
    # 3. INPUT VALIDATION & PHYSICAL BOUNDARY RULES
    # =========================================================
    add_heading_1("3. Input Validation Architecture & Thermodynamic Limits")
    add_para(
        "Input validation operates as a dedicated filter prior to any arithmetic execution. This decoupled design ensures that malformed inputs, malicious injection strings, and physically impossible temperatures are rejected with informative, user-friendly guidance."
    )
    add_heading_2("3.1 Thermodynamic Boundaries (Absolute Zero)")
    add_para(
        "The third law of thermodynamics defines Absolute Zero as the theoretical temperature at which the entropy of a system reaches its minimum possible value. No physical system can possess a temperature below Absolute Zero. The validation logic enforces the following lower bounds:"
    )
    add_bullet(" Values strictly below -273.15°C are rejected with a physical limit violation alert.", "Celsius Lower Bound: -273.15°C.")
    add_bullet(" Values strictly below -459.67°F are rejected with an explicit threshold warning.", "Fahrenheit Lower Bound: -459.67°F.")
    add_bullet(" Values strictly below 0.0 K are rejected immediately, preventing negative Kelvin values.", "Kelvin Lower Bound: 0.0 K.")

    add_heading_2("3.2 Syntactic & Numerical Rules")
    add_bullet(" Empty inputs and pure whitespace strings are detected and suppressed with an instructional prompt.", "Empty String Handling:")
    add_bullet(" Inputs are evaluated against the strict numerical regex: ^[+-]?((\\d+(\\.\\d*)?)|(\\.\\d+))([eE][+-]?\\d+)?$, correctly allowing scientific exponential notations (e.g. 1e3) while rejecting non-numeric tokens (e.g. 'abc', '12.34.56').", "Regular Expression Syntax Check:")
    add_bullet(" The parsed number is verified with Number.isFinite() to filter out Infinity and NaN results.", "Finiteness Assertion:")

    # =========================================================
    # 4. VISUAL UML SYSTEM DESIGN (RESOLVING MILESTONE 1 FEEDBACK)
    # =========================================================
    add_heading_1("4. Visual UML System Design & Architectural Models")
    add_para(
        "In accordance with specific faculty recommendations from the Milestone 1 evaluation, all UML system models have been formalized and visually rendered as high-resolution graphic diagrams. These models define the system use cases, static class hierarchies, and dynamic runtime sequence interactions."
    )

    add_heading_2("4.1 Use Case Diagram")
    add_para(
        "The Use Case Diagram depicts the system boundary and the interactions available to the end-user. The primary use cases include input specification, automated input validation with absolute zero boundary checks, high-precision conversion calculation, dynamic mathematical explanation rendering, and history persistence."
    )
    add_image_figure('diagrams/use_case_diagram.png', "Figure 1: Visual UML Use Case Diagram for the Temperature Conversion System")

    add_heading_2("4.2 UML Class Diagram")
    add_para(
        "The static structure of the prototype is organized into specialized, cohesive classes adhering to the Single Responsibility Principle (SRP):"
    )
    add_bullet(" Encapsulates numerical parsing, regex matching, and thermodynamic limit checks.", "InputValidator:")
    add_bullet(" Implements the mathematical formulas, IEEE 754 stabilization, and conversion routing.", "TemperatureConverter:")
    add_bullet(" Generates human-readable, step-by-step mathematical derivations and environmental state classifications.", "FormulaExplainer:")
    add_bullet(" Manages JSON serialization and browser persistence with automated quota handling and memory fallback.", "StorageManager:")
    add_bullet(" Handles DOM manipulation, event listeners, user feedback, and rendering.", "UIController:")
    add_image_figure('diagrams/class_diagram.png', "Figure 2: Visual UML Class Diagram illustrating Modular Object Decomposition")

    add_heading_2("4.3 UML Sequence Diagram")
    add_para(
        "The Sequence Diagram traces the chronological message exchange when a user enters a temperature value and triggers a conversion. It illustrates synchronous validation, scientific execution, derivation generation, and fault-tolerant storage logging."
    )
    add_image_figure('diagrams/sequence_diagram.png', "Figure 3: Visual UML Sequence Diagram for Conversion Execution and State Persistence")

    add_heading_2("4.4 Component Architecture")
    add_para(
        "The web application follows a clean 3-tier client-side architecture comprising the Presentation Layer, the Business Logic & Validation Layer, and the Persistence Layer."
    )
    add_image_figure('diagrams/system_architecture.png', "Figure 4: Client-Side Layered Component Architecture and Resiliency Infrastructure")

    # =========================================================
    # 5. RISK ASSESSMENT & MITIGATION MATRIX (RESOLVING MILESTONE 1 FEEDBACK)
    # =========================================================
    add_heading_1("5. Comprehensive Risk Assessment & Technical Mitigation")
    add_para(
        "Addressing the second key recommendation from the Milestone 1 review, this section presents a comprehensive risk assessment covering client-side web browser constraints, storage quotas, floating-point behaviors, and mitigation strategies implemented in Milestone 2."
    )

    # Risk Table
    risk_tbl = doc.add_table(rows=5, cols=4)
    risk_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_headers = ["Identified Risk", "Severity / Likelihood", "Potential Failure Mode", "Architectural Mitigation"]
    r_rows = [
        (
            "Browser LocalStorage Quotas & Privacy Restrictions",
            "Medium / High",
            "In private browsing mode (e.g. Safari Private Mode, Chrome Incognito) or under strict privacy sandboxes, window.localStorage throws a QuotaExceededError or DOMException: Access Denied.",
            "StorageManager implements proactive feature detection with try/catch wrapping. When localStorage is blocked or unavailable, it automatically switches to an in-memory session array without throwing unhandled exceptions."
        ),
        (
            "IEEE 754 Floating-Point Representation Errors",
            "Medium / Medium",
            "Fractional conversions (such as 5/9 or 9/5) can produce repeating binary representations and rounding artifacts (e.g. 99.99999999999999 instead of 100.0).",
            "Arithmetic uses scaled factor rounding with Number.EPSILON. The UI provides a selectable precision selector (2, 4, 6, 8 decimal places or raw) allowing users to select appropriate scientific precision."
        ),
        (
            "Cross-Browser DOM & CSS Variable Compatibility",
            "Low / Low",
            "Legacy mobile or outdated browser engines may lack support for modern CSS Grid, Flexbox, or ES6 class syntax.",
            "CSS styling adheres to W3C standard flexbox and responsive media queries. JavaScript utilizes standard ES6 classes supported across 98%+ of global web clients (Chrome 60+, Safari 11+, Firefox 55+, Edge 79+)."
        ),
        (
            "Invalid Input & Astronomical Number Injection",
            "Medium / Low",
            "Users entering non-numeric characters, exponential overflows (e.g. 1e308), or negative absolute temperatures.",
            "InputValidator executes multi-level boundary checking, enforcing finite number validation, scientific exponential parsing, and strict rejection of temperatures below Absolute Zero."
        )
    ]
    for c_idx, h_text in enumerate(r_headers):
        cell = risk_tbl.rows[0].cells[c_idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9)

    for r_idx, (col0, col1, col2, col3) in enumerate(r_rows):
        row = risk_tbl.rows[r_idx + 1]
        bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate([col0, col1, col2, col3]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 70, 70)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================
    # 6. PROTOTYPE USER INTERFACE & IMPLEMENTATION DETAILS
    # =========================================================
    add_heading_1("6. Prototype Implementation & UI Feature Set")
    add_para(
        "The Milestone 2 prototype has been developed and organized into clean, modular files located in the project's prototype directory:"
    )
    add_bullet(" Semantic HTML5 markup containing accessible input controls, ARIA live-regions, responsive grid containers, preset chips, and tabular history display.", "index.html:")
    add_bullet(" Modern design system using CSS variables, custom typography, subtle gradients, focus states, responsive media breakpoints, and color-coded status badges.", "styles.css:")
    add_bullet(" Clean, modular JavaScript containing the InputValidator, TemperatureConverter, FormulaExplainer, StorageManager, and UIController.", "converter.js:")
    add_bullet(" Standalone unit test suite that executes 30 distinct validation and conversion test assertions with zero dependencies.", "test_runner.js & tests.html:")

    add_heading_2("6.1 Core Prototype Features")
    add_bullet(" Instant calculation as the user types, alongside an explicit 'Convert Temperature' trigger button.", "Real-Time & Manual Conversion:")
    add_bullet(" A dedicated swap button instantly inverts the 'From' and 'To' scales and updates calculations automatically.", "Scale Inversion (Swap):")
    add_bullet(" Dynamic step-by-step breakdown explaining the exact mathematical derivation applied to the input.", "Mathematical Derivation Display:")
    add_bullet(" A contextual environmental badge classifying the entered temperature (e.g. Freezing point, Normal Body Temperature, Boiling point, Absolute Zero).", "Physical State Indicator:")
    add_bullet(" One-click clipboard copy functionality with visual confirmation feedback.", "Clipboard Copying:")
    add_bullet(" One-click buttons to load standard physical benchmarks (0°C, 100°C, 37°C, 0 K).", "Quick Preset Chips:")
    add_bullet(" Real-time session history displaying timestamp, input value, scales, output value, and conversion status.", "Tabular History Log:")

    add_heading_2("6.2 Prototype Interface Verification")
    add_para(
        "Figure 5 displays the live interactive prototype executing in a Chromium browser environment. The interface demonstrates the conversion of 100°C to 212.0000°F with dynamic derivation breakdown, boiling point thermal classification, and immediate tabular history logging."
    )
    add_image_figure('prototype/prototype_test_100C.png', "Figure 5: High-Fidelity Milestone 2 Interactive Prototype in Operation (Converting 100°C to 212°F)")

    # =========================================================
    # 7. QUALITY ASSURANCE, VERIFICATION & TEST RESULTS
    # =========================================================
    add_heading_1("7. Quality Assurance & Test Verification Suite")
    add_para(
        "To verify system correctness and satisfy Milestone 2 deliverables, an automated test suite was executed against all conversion and validation modules. A total of 30 test cases were executed, covering syntactic validation, physical limits, benchmark conversions, round-trip symmetry, and precision formatting."
    )

    # Test Results Table
    test_tbl = doc.add_table(rows=9, cols=4)
    test_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Test Category", "Test Condition / Input", "Expected Outcome", "Status"]
    t_rows = [
        ("Input Validation", "Empty string ('') and whitespace ('   ')", "Rejected with validation error", "PASSED (100%)"),
        ("Input Validation", "Non-numeric characters ('abc', '12.34.56')", "Rejected with format error", "PASSED (100%)"),
        ("Boundary Limits", "Absolute Zero boundary: -273.15°C and 0 K", "Accepted as valid physical boundary", "PASSED (100%)"),
        ("Boundary Limits", "Sub-Absolute Zero: -273.16°C and -0.001 K", "Rejected with physical limit violation", "PASSED (100%)"),
        ("Conversion Accuracy", "Freezing point: 0°C -> 32°F and 273.15 K", "Exact match: 32°F and 273.15 K", "PASSED (100%)"),
        ("Conversion Accuracy", "Boiling point: 100°C -> 212°F and 373.15 K", "Exact match: 212°F and 373.15 K", "PASSED (100%)"),
        ("Conversion Accuracy", "Scale convergence point: -40°C", "Exact match: -40°F", "PASSED (100%)"),
        ("Round-Trip Symmetry", "C -> F -> C and C -> K -> C round trips", "Difference <= 1e-9 (zero loss)", "PASSED (100%)")
    ]
    for c_idx, h_text in enumerate(t_headers):
        cell = test_tbl.rows[0].cells[c_idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9)

    for r_idx, (col0, col1, col2, col3) in enumerate(t_rows):
        row = test_tbl.rows[r_idx + 1]
        bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate([col0, col1, col2, col3]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if c_idx == 3:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x15, 0x80, 0x3D) # Emerald green

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================
    # 8. CONCLUSION & NEXT STEPS (MILESTONE 3 ROADMAP)
    # =========================================================
    add_heading_1("8. Conclusion & Roadmap for Milestone 3")
    add_para(
        "Milestone 2 has successfully realized the core conversion and validation modules, constructed an interactive web prototype, and comprehensively fulfilled the recommendations from the Milestone 1 review. The system is accurate, thermodynamically consistent, responsive, and resilient against client-side browser storage constraints."
    )
    add_para(
        "Looking forward to Milestone 3, the following technical extensions are planned:"
    )
    add_bullet(" Integration of Rankine (°R), Delisle (°De), and Réaumur (°Ré) scales into the modular conversion engine.", "1. Additional Scientific Scales:")
    add_bullet(" Implementation of client-side CSV and JSON export routines for the conversion history table.", "2. Data Export Utilities:")
    add_bullet(" Integration of lightweight SVG/Canvas charts to visualize temperature comparisons across scales dynamically.", "3. Interactive Temperature Comparison Visualizer:")
    add_bullet(" Packaging as a Progressive Web App (PWA) with a Service Worker for offline capability.", "4. Offline PWA Support:")

    # Save Document
    out_filename = "Temperature_Conversion_System_Milestone_2_Report.docx"
    doc.save(out_filename)
    print(f"Report document successfully generated: {out_filename}")

if __name__ == "__main__":
    create_report()
